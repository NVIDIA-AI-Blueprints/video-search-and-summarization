# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Focused checks for fixed-load latency captured during a max-live search."""

import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from latency_tracker import LatencyTracker
from live_streams_benchmark import (
    LiveStreamsBenchmark,
    next_ramp_target,
    summarize_latency_plateau,
)


class LiveStreamPlateauTest(unittest.TestCase):
    def test_ramp_lands_on_requested_counts(self):
        checkpoints = [1, 16, 32, 64, 128]
        self.assertEqual(next_ramp_target(1, 5, checkpoints), 6)
        self.assertEqual(next_ramp_target(11, 5, checkpoints), 16)
        self.assertEqual(next_ramp_target(31, 5, checkpoints), 32)
        self.assertEqual(next_ramp_target(32, 5, checkpoints), 37)
        seen = [1]
        while seen[-1] < 128:
            seen.append(next_ramp_target(seen[-1], 5, checkpoints))
        self.assertTrue(set(checkpoints).issubset(seen))

    def test_plateau_has_full_window_distribution_and_per_stream_counts(self):
        tracker = LatencyTracker()
        for stream_id, values in {"a": [1, 2], "b": [3, 4]}.items():
            for value in values:
                tracker.record_latency(value, stream_id)

        result = summarize_latency_plateau(
            tracker, ["a", "b"], 300.0, 300, 0, 10.0
        )

        self.assertTrue(result["valid"])
        self.assertEqual(result["stream_count"], 2)
        self.assertEqual(result["total_measurements"], 4)
        self.assertAlmostEqual(result["p95_latency"], 3.85)
        self.assertEqual(result["per_stream_stats"]["a"]["total_measurements"], 2)
        self.assertEqual(result["latency_history"], {"a": [1, 2], "b": [3, 4]})

    def test_plateau_rejects_missing_stream_or_drops(self):
        tracker = LatencyTracker()
        tracker.record_latency(1, "a")
        result = summarize_latency_plateau(
            tracker, ["a", "b"], 300.0, 300, 1, 10.0
        )
        self.assertFalse(result["valid"])
        self.assertEqual(result["fresh_streams"], 1)
        self.assertEqual(result["dropped_chunks"], 1)

    def test_plateau_rejects_stale_interval_even_with_final_samples(self):
        tracker = LatencyTracker()
        tracker.record_latency(1, "a")
        result = summarize_latency_plateau(
            tracker, ["a"], 300.0, 300, 0, 10.0, all_windows_fresh=False
        )
        self.assertFalse(result["valid"])
        self.assertIn("stale_window", result["invalid_reasons"])

    def test_mixed_timestamp_and_missing_timestamp_chunks_fail_plateau(self):
        benchmark = LiveStreamsBenchmark("http://localhost")
        valid = {"chunk_id": 1, "media_info": {
            "type": "timestamp",
            "end_timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%fZ"),
        }}
        missing = {"chunk_id": 2, "chunk_responses": [{"processing_latency_s": 1.0}]}
        malformed = {"chunk_id": 3, "media_info": {
            "type": "timestamp", "end_timestamp": "invalid"
        }}
        benchmark._record_result_latency("stream-1", valid, "ntp_timestamp", True)
        benchmark._record_result_latency("stream-1", missing, "ntp_timestamp", True)
        benchmark._record_result_latency("stream-1", malformed, "ntp_timestamp", True)
        result = summarize_latency_plateau(
            benchmark.latency_tracker, ["stream-1"], 300.0, 300, 0, 10.0,
            missing_latency_events=benchmark._get_missing_latency_count(["stream-1"]),
        )
        self.assertEqual(result["total_measurements"], 1)
        self.assertEqual(result["missing_latency_events"], 2)
        self.assertFalse(result["valid"])
        self.assertIn("missing_latency_source", result["invalid_reasons"])
        benchmark._reset_stream_integrity_tracking()
        self.assertEqual(benchmark._get_missing_latency_count(["stream-1"]), 0)

    def test_excel_report_exposes_fixed_load_latency(self):
        with tempfile.TemporaryDirectory() as root:
            results = Path(root)
            case_id = "case"
            (results / case_id).mkdir()
            (results / "execution_summary.json").write_text(json.dumps({
                "test_cases": [{"test_case_id": case_id}],
            }))
            (results / case_id / "max_live_streams_results.json").write_text(json.dumps({
                "latency_plateau_counts": [16, 32],
                "latency_plateaus": [{
                    "stream_count": 16,
                    "duration_seconds": 300,
                    "total_measurements": 32,
                    "p95_latency": 2.5,
                    "valid": False,
                    "invalid_reasons": ["stale_window"],
                    "latency_history": {"a": [2.5]},
                    "per_stream_stats": {"a": {"total_measurements": 1}},
                }],
            }))
            report = results / "report.xlsx"
            LiveStreamsBenchmark("http://localhost").analyze_results(str(results), str(report))
            book = load_workbook(report, read_only=True)
            self.assertIn("Fixed_Load_Latency", book.sheetnames)
            rows = list(book["Fixed_Load_Latency"].values)
            self.assertEqual(rows[1][rows[0].index("stream_count")], 16)
            self.assertEqual(rows[1][rows[0].index("invalid_reasons")], "stale_window")
            self.assertEqual(rows[2][rows[0].index("stream_count")], 32)
            self.assertEqual(rows[2][rows[0].index("invalid_reasons")], "unreached")
            self.assertIn("Fixed_Load_Per_Stream", book.sheetnames)
            stream_rows = list(book["Fixed_Load_Per_Stream"].values)
            self.assertEqual(stream_rows[1][stream_rows[0].index("stream_id")], "a")
            self.assertEqual(stream_rows[1][stream_rows[0].index("total_measurements")], 1)

    def test_execute_counts_incomplete_fixed_load_as_failure(self):
        benchmark = LiveStreamsBenchmark("http://localhost")
        config = {"global": {}, "test_scenarios": {"case": {"benchmark_mode": "max_live_streams", "videos": [{
            "rtsp_url": "rtsp://example/test", "chunk_sizes": [10], "latency_threshold_seconds": 10,
        }]}}}
        with tempfile.TemporaryDirectory() as root, patch.object(
            benchmark, "setup_scenario_directory", return_value=root
        ), patch.object(benchmark, "parse_global_config", return_value={}), patch.object(
            benchmark, "parse_benchmark_config", return_value=config["test_scenarios"]["case"]
        ), patch.object(benchmark, "get_available_models", return_value="model"), patch.object(
            benchmark, "_execute_live_streams_test_case",
            return_value={"test_case_id": "case", "success": True, "fixed_load_latency_complete": False},
        ):
            result = benchmark.execute(config, "case")
        self.assertEqual(result["successful_test_cases"], 0)
        self.assertEqual(result["failed_test_cases"], 1)

    def test_invalid_checkpoints_fail_before_starting_monitors(self):
        benchmark = LiveStreamsBenchmark("http://localhost")
        with tempfile.TemporaryDirectory() as root, patch.object(
            benchmark, "start_gpu_monitoring"
        ) as start_gpu:
            with self.assertRaisesRegex(ValueError, "latency_plateau_counts"):
                benchmark._execute_live_streams_test_case(
                    "case",
                    {
                        "initial_stream_count": 5,
                        "latency_plateau_counts": [1, 16],
                        "latency_threshold_seconds": 10,
                    },
                    10,
                    {"backend_type": "rtvi_vlm"},
                    "model",
                    root,
                )
            start_gpu.assert_not_called()

    def test_fixed_load_requires_effective_ntp_source_before_starting_monitors(self):
        benchmark = LiveStreamsBenchmark("http://localhost")
        with tempfile.TemporaryDirectory() as root, patch.object(
            benchmark, "start_gpu_monitoring"
        ) as start_gpu:
            with self.assertRaisesRegex(ValueError, "NTP timestamp"):
                benchmark._execute_live_streams_test_case(
                    "case",
                    {
                        "initial_stream_count": 1,
                        "latency_plateau_counts": [1],
                        "latency_measurement_source": "processing_latency",
                        "latency_threshold_seconds": 10,
                    },
                    10,
                    {"backend_type": "rtvi_vlm", "latency_measurement_source": "ntp_timestamp"},
                    "model",
                    root,
                )
            start_gpu.assert_not_called()


if __name__ == "__main__":
    unittest.main()
