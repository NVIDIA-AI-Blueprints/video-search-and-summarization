# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Advisory live capacity derived from matching sustained file workloads."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from file_burst_benchmark import FileBurstBenchmark, estimate_live_streams
from openpyxl import load_workbook


class FileLiveEstimateTest(unittest.TestCase):
    def setUp(self):
        self.video = {"file_duration_seconds": 10, "live_chunk_duration_seconds": 10}
        self.row = {
            "concurrency_level": 64,
            "completed_files": 2310,
            "failed_files": 0,
            "steady_state_duration_seconds": 300,
            "e2e_latency_seconds": 300,
            "throughput_files_per_second": 7.7,
            "p95_latency": 8.25,
        }

    def test_best_eligible_throughput_and_margin(self):
        rows = [
            dict(self.row, throughput_files_per_second=7.13),
            self.row,
            dict(self.row, throughput_files_per_second=9, p95_latency=11),
            dict(self.row, throughput_files_per_second=10, failed_files=1),
        ]
        result = estimate_live_streams(rows, self.video, 0)
        self.assertEqual(result["estimated_live_streams"], 77)
        self.assertEqual(result["suggested_initial_stream_count"], 72)
        self.assertEqual(result["source_concurrency_level"], 64)
        self.assertFalse(result["validated_live_capacity"])

    def test_opt_in_and_matching_single_chunk_required(self):
        self.assertIsNone(estimate_live_streams([self.row], {}, 0))
        for video, chunk in [
            ({"live_chunk_duration_seconds": 10}, 0),
            (dict(self.video, file_duration_seconds=120), 0),
            (self.video, 5),
            (dict(self.video, live_chunk_duration_seconds=True), 0),
            (dict(self.video, file_duration_seconds=float("nan")), 0),
        ]:
            with self.subTest(video=video, chunk=chunk), self.assertRaises(ValueError):
                estimate_live_streams([self.row], video, chunk)
        self.assertEqual(
            estimate_live_streams([self.row], self.video, 10)["status"], "estimated"
        )

    def test_invalid_or_non_sustained_measurements_are_not_estimates(self):
        for field, value in [
            ("failed_files", None),
            ("completed_files", 0),
            ("steady_state_duration_seconds", 0),
            ("e2e_latency_seconds", 5),
            ("p95_latency", float("nan")),
            ("p95_latency", -1),
            ("throughput_files_per_second", float("inf")),
            ("throughput_files_per_second", 0),
        ]:
            with self.subTest(field=field, value=value):
                result = estimate_live_streams(
                    [dict(self.row, **{field: value})], self.video, 0
                )
                self.assertEqual(result["status"], "unavailable")
                self.assertNotIn("estimated_live_streams", result)
        self.assertEqual(
            estimate_live_streams([], self.video, 0)["status"], "unavailable"
        )

    def test_floor_and_minimum_start(self):
        result = estimate_live_streams(
            [dict(self.row, throughput_files_per_second=0.19, p95_latency=10)],
            self.video,
            0,
        )
        self.assertEqual(result["estimated_live_streams"], 1)
        self.assertEqual(result["suggested_initial_stream_count"], 1)

    def test_producer_keeps_precision_for_live_estimates(self):
        benchmark = FileBurstBenchmark("http://localhost")
        video = {
            "filepath": "clip.mp4",
            "file_duration_seconds": 100,
            "live_chunk_duration_seconds": 100,
            "steady_state_duration_seconds": 300,
            "reuse_uploaded_files": False,
        }
        for p95, expected_status in ((90.0, "estimated"), (100.004, "unavailable")):
            with (
                self.subTest(p95=p95),
                tempfile.TemporaryDirectory() as root,
                patch("file_burst_benchmark.time") as clock,
                patch.object(benchmark, "start_gpu_monitoring"),
                patch.object(benchmark, "stop_gpu_monitoring"),
                patch.object(benchmark, "process_gpu_stats", return_value={}),
                patch.object(
                    benchmark, "_process_single_file_burst",
                    return_value={"success": True, "processing_time": p95},
                ),
            ):
                clock.time.side_effect = [0.0, 1000.0, 1000.0]
                row = benchmark._test_single_concurrency_level(
                    "precision", video, 0, 96, {}, "model", root
                )
                estimate = estimate_live_streams([row], video, 0)
                self.assertEqual(estimate["status"], expected_status)
                self.assertEqual(row["throughput_files_per_second"], 0.096)
                self.assertAlmostEqual(row["p95_latency"], p95)
                if expected_status == "estimated":
                    self.assertEqual(estimate["estimated_live_streams"], 9)
                    self.assertEqual(estimate["suggested_initial_stream_count"], 4)

    def test_iteration_persists_estimate_and_excel_keeps_failed_case_visible(self):
        benchmark = FileBurstBenchmark("http://localhost")
        video = dict(
            self.video,
            filepath="clip.mp4",
            concurrency_levels=[64],
            target_latency_seconds=0,
        )
        with tempfile.TemporaryDirectory() as root:
            case = Path(root) / "case"
            case.mkdir()
            # A failed overload probe must not hide a valid lower-concurrency estimate.
            rows = [self.row, dict(self.row, concurrency_level=128, failed_files=1)]
            with (
                patch.object(
                    benchmark, "_test_single_concurrency_level", side_effect=rows
                ),
                patch("file_burst_benchmark.time.sleep"),
            ):
                result = benchmark._execute_file_burst_iteration(
                    1,
                    dict(video, concurrency_levels=[64, 128]),
                    0,
                    {},
                    "model",
                    str(case),
                )
            saved = json.loads((case / "file_burst_results.json").read_text())
            self.assertEqual(
                saved["live_stream_estimate"]["estimated_live_streams"], 77
            )
            self.assertFalse(saved["success"])
            (case / "test_case_summary.json").write_text(
                json.dumps(
                    {
                        "video_file": "clip.mp4",
                        "iteration_results": [result],
                    }
                )
            )
            (Path(root) / "execution_summary.json").write_text(
                json.dumps(
                    {
                        "test_cases": [{"test_case_id": "case", "success": False}],
                    }
                )
            )
            with patch.object(
                benchmark, "get_gpu_info_dataframe", return_value=pd.DataFrame()
            ):
                benchmark.analyze_results(root, str(Path(root) / "report.xlsx"))
            workbook = load_workbook(Path(root) / "report.xlsx")
            sheet = workbook["Live_Stream_Estimates"]
            values = list(sheet.values)
            row = dict(zip(values[0], values[1]))
            self.assertEqual(row["estimated_live_streams"], 77)
            self.assertEqual(row["iteration"], 1)
            self.assertFalse(row["iteration_success"])
            self.assertFalse(row["validated_live_capacity"])
            workbook.close()

    def test_old_report_has_no_estimate_sheet(self):
        benchmark = FileBurstBenchmark("http://localhost")
        with tempfile.TemporaryDirectory() as root:
            case = Path(root) / "case"
            case.mkdir()
            (case / "test_case_summary.json").write_text(
                json.dumps({"iteration_results": []})
            )
            (Path(root) / "execution_summary.json").write_text(
                json.dumps(
                    {
                        "test_cases": [{"test_case_id": "case", "success": True}],
                    }
                )
            )
            with patch.object(
                benchmark, "get_gpu_info_dataframe", return_value=pd.DataFrame()
            ):
                benchmark.analyze_results(root, str(Path(root) / "report.xlsx"))
            workbook = load_workbook(Path(root) / "report.xlsx")
            self.assertNotIn("Live_Stream_Estimates", workbook.sheetnames)
            workbook.close()

    def test_config_validation_is_opt_in_and_rejects_embedding_backend(self):
        benchmark = FileBurstBenchmark("http://localhost")
        video = {"filepath": "clip.mp4", "chunk_sizes": [0], "concurrency_levels": [1]}
        config = {"benchmark_mode": "file_burst", "videos": [video]}
        self.assertEqual(
            benchmark.parse_benchmark_config(config, {})["videos"], [video]
        )
        video.update(self.video)
        with self.assertRaisesRegex(ValueError, "rtvi_vlm backend"):
            benchmark.parse_benchmark_config(config, {"backend_type": "rtvi_embed"})
        video["file_duration_seconds"] = 120
        with self.assertRaisesRegex(ValueError, "equal"):
            benchmark.parse_benchmark_config(config, {})


if __name__ == "__main__":
    unittest.main()
