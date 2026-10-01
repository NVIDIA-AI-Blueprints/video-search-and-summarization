# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Cross-platform BCD report labels must describe each platform's workload."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import yaml
from openpyxl import Workbook

from generate_bcd_3_3_configs import HERE
from generate_perf_xlsx import write_machine_config_tab, write_summary_tab


class BCDReportLabelsTest(unittest.TestCase):
    def test_machine_config_uses_each_platform_frame_budget(self):
        configs = {
            label: yaml.safe_load(
                (HERE / f"rtvi_vlm_bcd_3_3_{platform}_config.yaml").read_text()
            )
            for label, platform in (("Orin", "agx_orin"), ("RTX", "rtx_pro_6000_se"))
        }
        book = Workbook()
        write_machine_config_tab(
            book, configs, {label: Path("/tmp") for label in configs}
        )
        rows = list(book["Machine Config"].values)
        scenario = "max_live_streams_test_100_token_2k"
        values = {row[1]: row for row in rows if row[0] == scenario}
        self.assertEqual(values["Orin"][3], 5)
        self.assertEqual(values["RTX"][3], 10)
        self.assertEqual(values["Orin"][4], values["RTX"][4])

    def test_machine_config_includes_platform_only_scenarios(self):
        configs = {
            "H100": yaml.safe_load((HERE / "rtvi_vlm_config_h100.yaml").read_text()),
            "Orin": yaml.safe_load(
                (HERE / "rtvi_vlm_bcd_3_3_agx_orin_config.yaml").read_text()
            ),
        }
        book = Workbook()
        write_machine_config_tab(
            book, configs, {name: Path("/tmp") for name in configs}
        )
        rows = list(book["Machine Config"].values)
        self.assertIn(
            ("max_live_streams_test_1_token", "H100"), [row[:2] for row in rows]
        )
        self.assertIn(
            ("max_live_streams_test_1_token_4k", "Orin"), [row[:2] for row in rows]
        )

    def test_max_stream_summary_keeps_vision_tiers_distinct(self):
        config = yaml.safe_load(
            (HERE / "rtvi_vlm_bcd_3_3_agx_orin_config.yaml").read_text()
        )
        with tempfile.TemporaryDirectory() as root:
            report = Path(root)
            for tier, streams in (("2k", 8), ("4k", 4)):
                case = report / f"max_live_streams_test_100_token_{tier}" / "case"
                case.mkdir(parents=True)
                (case / "max_live_streams_results.json").write_text(
                    json.dumps({"success": True, "max_sustainable_streams": streams})
                )
            book = Workbook()
            write_summary_tab(
                book,
                {"Orin": report},
                {
                    "release": "3.3",
                    "model": "Cosmos3-Edge",
                    "precision": "BF16",
                    "engine": "vLLM",
                    "isl_text": 0,
                },
                {"Orin": config},
            )
        labels = [row[0] for row in book["Summary"].values if isinstance(row[0], str)]
        self.assertTrue(any("VT=2000" in label for label in labels))
        self.assertTrue(any("VT=4000" in label for label in labels))

    def test_max_stream_chart_labels_each_platform_frame_budget(self):
        from plot_perf_reports import (
            plot_max_streams_2k,
            plot_max_streams_2k_vs_8k,
            plot_max_streams_8k,
        )

        configs = {
            label: yaml.safe_load(
                (HERE / f"rtvi_vlm_bcd_3_3_{platform}_config.yaml").read_text()
            )
            for label, platform in (("Orin", "agx_orin"), ("RTX", "rtx_pro_6000_se"))
        }
        configs["H100"] = yaml.safe_load(
            (HERE / "rtvi_vlm_config_h100.yaml").read_text()
        )
        configs["H100"]["test_scenarios"]["max_live_streams_test_100_token_448"][
            "videos"
        ][0]["generate_captions_params"][
            "num_frames_per_second_or_fixed_frames_chunk"
        ] = 64
        with tempfile.TemporaryDirectory() as root:
            reports = {}
            for label in configs:
                report = Path(root) / label
                scenarios = (
                    [
                        "max_live_streams_test_100_token",
                        "max_live_streams_test_100_token_448",
                    ]
                    if label == "H100"
                    else [
                        f"max_live_streams_test_100_token_{tier}"
                        for tier in ("2k", "4k", "8k")
                    ]
                )
                for scenario in scenarios:
                    case = report / scenario / "case"
                    case.mkdir(parents=True)
                    (case / "max_live_streams_results.json").write_text(
                        json.dumps(
                            {"max_sustainable_streams": 2, "last_stable_p95": 1.0}
                        )
                    )
                reports[label] = report
            with patch("plot_perf_reports.plt.close"):
                plot_max_streams_2k(reports, configs, Path(root))
                from matplotlib import pyplot as plt

                labels = [
                    tick.get_text() for tick in plt.gcf().axes[1].get_xticklabels()
                ]
                self.assertIn("ORIN\n5 frames", labels)
                self.assertIn("RTX\n10 frames", labels)
                self.assertIn("H100\n30 frames", labels)
                plot_max_streams_8k(reports, configs, Path(root))
                labels = [
                    tick.get_text() for tick in plt.gcf().axes[1].get_xticklabels()
                ]
                self.assertIn("H100\n64 frames", labels)
                plot_max_streams_2k_vs_8k(reports, configs, Path(root))
                labels = [
                    tick.get_text() for tick in plt.gcf().axes[1].get_xticklabels()
                ]
                self.assertIn("ORIN\n2K/4K/8K: 5/10/20 frames", labels)
                self.assertIn("RTX\n2K/4K/8K: 10/20/40 frames", labels)
                self.assertIn("H100\n2K/4K/8K: 30/–/64 frames", labels)
                plt.close("all")


if __name__ == "__main__":
    unittest.main()
